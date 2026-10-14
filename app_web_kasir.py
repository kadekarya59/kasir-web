import os
from datetime import datetime
import streamlit as st
import openpyxl
from openpyxl import Workbook
import qrcode

from reportlab.lib.pagesizes import A6
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

DB_FILE = "database_kasir.xlsx"

# ---------------------------------------------------------
# DATABASE FUNCTIONS
# ---------------------------------------------------------
def init_database():
    """Membuat database spreadsheet jika belum ada"""
    if not os.path.exists(DB_FILE):
        wb = Workbook()
        
        # Sheet Produk
        ws_produk = wb.active
        ws_produk.title = "Produk"
        ws_produk.append(["Kode Produk", "Nama Produk", "Harga (Rp)", "Stok", "Kategori", "Icon"])
        
        produk_awal = [
            ["P001", "Kopi Susu Gula Aren", 18000, 50, "Minuman", "☕"],
            ["P002", "Roti Bakar Cokelat", 15000, 30, "Makanan", "🍞"],
            ["P003", "Teh Manis Dingin", 6000, 100, "Minuman", "🧃"],
            ["P004", "Mie Goreng Spesial", 22000, 25, "Makanan", "🍜"],
            ["P005", "Air Mineral 600ml", 4000, 80, "Minuman", "🍼"],
        ]
        for p in produk_awal:
            ws_produk.append(p)
            
        # Sheet Transaksi
        ws_transaksi = wb.create_sheet(title="Transaksi")
        ws_transaksi.append(["No Nota", "Tanggal", "Kode Produk", "Nama Produk", "Qty", "Harga Satuan", "Total Harga", "Total Belanja", "Metode Bayar", "Tunai", "Kembalian"])
        
        wb.save(DB_FILE)

def load_produk():
    init_database()
    wb = openpyxl.load_workbook(DB_FILE)
    ws = wb["Produk"]
    produk_dict = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            produk_dict[str(row[0])] = {
                "nama": row[1],
                "harga": int(row[2]),
                "stok": int(row[3]),
                "kategori": row[4] if len(row) > 4 and row[4] else "Umum",
                "icon": row[5] if len(row) > 5 and row[5] else "📦"
            }
    return produk_dict

def tambah_atau_update_produk_db(kode, nama, harga, stok, kategori, icon):
    wb = openpyxl.load_workbook(DB_FILE)
    ws = wb["Produk"]
    found = False
    
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == kode:
            row[1].value = nama
            row[2].value = harga
            row[3].value = stok
            row[4].value = kategori
            row[5].value = icon
            found = True
            break
            
    if not found:
        ws.append([kode, nama, harga, stok, kategori, icon])
        
    wb.save(DB_FILE)

def hapus_produk_db(kode):
    wb = openpyxl.load_workbook(DB_FILE)
    ws = wb["Produk"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == kode:
            ws.delete_rows(row[0].row)
            break
    wb.save(DB_FILE)

def update_stok_dan_simpan_transaksi(no_nota, keranjang, total_bayar, metode_bayar, tunai, kembalian):
    wb = openpyxl.load_workbook(DB_FILE)
    ws_produk = wb["Produk"]
    ws_trans = wb["Transaksi"]
    tgl_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for item in keranjang:
        kode = item["kode"]
        qty = item["qty"]
        
        for row in ws_produk.iter_rows(min_row=2):
            if str(row[0].value) == kode:
                row[3].value = int(row[3].value) - qty
                break
                
        ws_trans.append([
            no_nota,
            tgl_str,
            kode,
            item["nama"],
            qty,
            item["harga"],
            item["subtotal"],
            total_bayar,
            metode_bayar,
            tunai,
            kembalian
        ])
        
    wb.save(DB_FILE)

def cari_nota_di_db(no_nota_cari):
    if not os.path.exists(DB_FILE):
        return None
    wb = openpyxl.load_workbook(DB_FILE)
    if "Transaksi" not in wb.sheetnames:
        return None
    ws_trans = wb["Transaksi"]
    items = []
    header_info = {}
    for row in ws_trans.iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).strip().upper() == str(no_nota_cari).strip().upper():
            if not header_info:
                header_info = {
                    "no_nota": row[0],
                    "tanggal": row[1],
                    "total_belanja": row[7],
                    "metode_bayar": row[8] if len(row) > 8 else "Cash / Tunai",
                    "tunai": row[9] if len(row) > 9 else row[7],
                    "kembalian": row[10] if len(row) > 10 else 0
                }
            items.append({
                "kode": row[2],
                "nama": row[3],
                "qty": row[4],
                "harga": row[5],
                "subtotal": row[6]
            })
    if items:
        return {"header": header_info, "items": items}
    return None

# ---------------------------------------------------------
# GENERATE PDF & QR CODE
# ---------------------------------------------------------
def generate_qr_code(text):
    qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_L, box_size=4, border=1)
    qr.add_data(text)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    filename = f"qr_{text}.png"
    img.save(filename)
    return filename

def buat_nota_pdf(no_nota, keranjang, total_bayar, metode_bayar, tunai, kembalian, tgl_str=None):
    filename = f"Nota_{no_nota}.pdf"
    doc = SimpleDocTemplate(filename, pagesize=A6, rightMargin=10, leftMargin=10, topMargin=10, bottomMargin=10)
    styles = getSampleStyleSheet()
    
    style_center = ParagraphStyle(name='Center', parent=styles['Normal'], alignment=1, fontSize=9)
    style_title = ParagraphStyle(name='TitleStyle', parent=styles['Heading2'], alignment=1, fontSize=12, leading=14)
    style_text = ParagraphStyle(name='TextStyle', parent=styles['Normal'], fontSize=8, leading=10)
    
    elements = []
    elements.append(Paragraph("<b>MINIMARKET MART DIGITAL</b>", style_title))
    elements.append(Paragraph("Jl. Merdeka No. 123, Indonesia", style_center))
    elements.append(Paragraph("Telp: 0812-3456-7890", style_center))
    elements.append(Spacer(1, 6))
    elements.append(Paragraph("-" * 45, style_center))
    
    if not tgl_str:
        tgl_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    elements.append(Paragraph(f"<b>Kode Nota : {no_nota}</b>", style_text))
    elements.append(Paragraph(f"Tanggal   : {tgl_str}", style_text))
    elements.append(Spacer(1, 5))
    
    table_data = [["Item", "Qty", "Harga", "Total"]]
    for item in keranjang:
        table_data.append([item["nama"], str(item["qty"]), f"{item['harga']:,}", f"{item['subtotal']:,}"])
        
    t = Table(table_data, colWidths=[95, 25, 45, 55])
    t.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
        ('TOPPADDING', (0,0), (-1,-1), 2),
        ('LINEBELOW', (0,0), (-1,0), 0.5, colors.black),
        ('ALIGN', (1,0), (-1,-1), 'RIGHT'),
    ]))
    elements.append(t)
    elements.
